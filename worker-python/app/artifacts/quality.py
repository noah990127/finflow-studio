import io
import json
import re
from docx import Document
from pptx import Presentation
from pypdf import PdfReader

from ..models import DeliverableRequest
from .models import ArtifactFormat, ArtifactQualityReport, QualityIssue


def evaluate_artifact(
    output_format: ArtifactFormat, request: DeliverableRequest, content: bytes
) -> ArtifactQualityReport:
    issues: list[QualityIssue] = []
    checks = {"non_empty": bool(content), "has_content": _has_content(request)}
    metrics: dict[str, int | float | str | bool] = {"bytes": len(content)}
    if not content:
        _error(issues, "empty_file", "成果文件为空", repair="重新执行渲染并检查输入内容")
    if not checks["has_content"]:
        _error(issues, "empty_content", "成果没有可用正文", repair="补充分析结论或上游资料")
    _evaluate_references(request, issues, checks, metrics)
    try:
        if output_format == "pptx":
            _evaluate_pptx(request, content, issues, checks, metrics)
        elif output_format == "docx":
            _evaluate_docx(content, issues, checks, metrics)
        elif output_format == "pdf":
            _evaluate_pdf(content, issues, checks, metrics)
        elif output_format == "html_slides":
            _evaluate_html(content, issues, checks, metrics)
        elif output_format == "financial_report":
            _evaluate_dashboard(content, issues, checks, metrics)
        elif output_format == "mermaid":
            _evaluate_mermaid(content, issues, checks, metrics)
        elif output_format == "excalidraw":
            _evaluate_excalidraw(content, issues, checks, metrics)
    except Exception as exception:
        _error(issues, "unreadable_artifact", "成果无法重新解析", detail=str(exception), repair="检查渲染器输出协议")
    error_count = sum(issue.severity == "error" for issue in issues)
    warning_count = sum(issue.severity == "warning" for issue in issues)
    score = max(0, 100 - error_count * 35 - warning_count * 8)
    return ArtifactQualityReport(
        format=output_format,
        score=score,
        passed=error_count == 0 and score >= 70,
        checks=checks,
        issues=issues,
        metrics=metrics,
    )


def _evaluate_references(request, issues, checks, metrics):
    refs = [ref for section in request.sections for ref in section.refs]
    unique = {(ref.ref_id, ref.source_name, ref.content_hash) for ref in refs}
    metrics["source_count"] = len(unique)
    checks["citations_present"] = not request.include_citations or bool(unique)
    if request.include_citations and not unique:
        _warning(issues, "missing_citations", "已开启引用，但成果没有绑定来源", repair="从上游节点传入已验证引用")
    incomplete = [ref.source_name for ref in refs if not ref.resource_id or ref.version <= 0 or not ref.content_hash]
    if incomplete:
        _warning(issues, "incomplete_provenance", f"有 {len(incomplete)} 条来源缺少完整快照信息", repair="补齐资源 ID、版本和内容哈希")


def _evaluate_pptx(request, content, issues, checks, metrics):
    deck = Presentation(io.BytesIO(content))
    metrics["slide_count"] = len(deck.slides)
    checks["package_readable"] = True
    if len(deck.slides) < 2:
        _error(issues, "too_few_slides", "演示文稿缺少正文页面", repair="至少生成封面和一页正文")
    empty = []
    for index, slide in enumerate(deck.slides, start=1):
        text = " ".join(shape.text for shape in slide.shapes if hasattr(shape, "text"))
        if index > 1 and len(text.strip()) < 8:
            empty.append(index)
        if len(text) > 520:
            _warning(issues, "slide_density", f"第 {index} 页文字过多", str(index), "拆页或改用图表/结构图")
    if empty:
        _error(issues, "empty_slides", "存在空白正文页：" + ", ".join(map(str, empty)), repair="删除空页或补充内容")
    headings = [_normalized_text(section.heading) for section in request.sections if section.heading]
    duplicate_headings = sorted({heading for heading in headings if headings.count(heading) > 1})
    if duplicate_headings:
        _warning(
            issues,
            "duplicate_slide_titles",
            "正文页存在重复标题：" + "、".join(duplicate_headings),
            repair="将每页标题改为该页的唯一结论",
        )
    core_messages = [_normalized_text(section.core_message) for section in request.sections if section.core_message]
    duplicate_core = sorted({message for message in core_messages if core_messages.count(message) > 1})
    if duplicate_core:
        _error(
            issues,
            "duplicate_core_messages",
            "多页重复表达同一核心观点",
            repair="合并重复页，或为各页补充不同的证据与含义",
        )


def _evaluate_docx(content, issues, checks, metrics):
    document = Document(io.BytesIO(content))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    headings = [paragraph.text for paragraph in document.paragraphs if paragraph.style.name.startswith("Heading")]
    metrics.update({"paragraph_count": len(document.paragraphs), "heading_count": len(headings), "text_chars": len(text)})
    checks["package_readable"] = True
    if len(text.strip()) < 40:
        _error(issues, "document_too_short", "正式文档正文过短", repair="补齐摘要、分析和结论")
    if not headings:
        _error(issues, "missing_headings", "正式文档没有章节结构", repair="生成文档级大纲和标题层级")
    if any(len(paragraph.text) > 1800 for paragraph in document.paragraphs):
        _warning(issues, "long_paragraph", "存在过长段落，连续阅读困难", repair="按论点拆分段落")


def _evaluate_pdf(content, issues, checks, metrics):
    if not content.startswith(b"%PDF"):
        _error(issues, "invalid_pdf_signature", "文件不是有效 PDF", repair="检查 PDF 渲染器")
        return
    reader = PdfReader(io.BytesIO(content))
    texts = [(page.extract_text() or "").strip() for page in reader.pages]
    metrics.update({"page_count": len(reader.pages), "text_chars": sum(map(len, texts))})
    checks.update({"package_readable": True, "searchable_text": any(texts)})
    if not any(texts):
        _error(issues, "pdf_not_searchable", "PDF 没有可搜索文字", repair="使用文本排版而不是整页图片")
    empty = [index for index, value in enumerate(texts, start=1) if not value]
    if empty:
        _warning(issues, "blank_pdf_pages", "PDF 存在空白页：" + ", ".join(map(str, empty)), repair="调整分页规则")


def _evaluate_html(content, issues, checks, metrics):
    html = content.decode("utf-8")
    slides = len(re.findall(r'<section class="slide', html))
    metrics.update({"slide_count": slides, "text_chars": len(re.sub(r"<[^>]+>", "", html))})
    checks.update({"doctype": html.lower().startswith("<!doctype html>"), "keyboard_navigation": "keydown" in html,
                   "responsive_viewport": 'name="viewport"' in html})
    if not all(checks[key] for key in ("doctype", "keyboard_navigation", "responsive_viewport")):
        _error(issues, "html_contract", "网页演示缺少文档声明、响应式视口或键盘导航", repair="使用 HTML 演示模板重新渲染")
    if slides < 2:
        _error(issues, "html_too_few_slides", "网页演示缺少正文页面", repair="至少生成封面和一页正文")
    if re.search(r"<(?:script|link|img)\b[^>]+(?:src|href)=[\"']https?://", html, re.IGNORECASE):
        _warning(issues, "external_dependency", "网页演示包含外部资源，离线时可能无法显示", repair="将必要资源打包进成果")


def _evaluate_dashboard(content, issues, checks, metrics):
    payload = json.loads(content)
    sections = payload.get("sections") if isinstance(payload.get("sections"), list) else []
    views = payload.get("views") if isinstance(payload.get("views"), list) else []
    metrics.update({"section_count": len(sections), "view_count": len(views)})
    checks["schema_valid"] = payload.get("renderer") == "finbtp-echarts-perspective" and bool(sections)
    if not checks["schema_valid"]:
        _error(issues, "dashboard_schema", "交互报告协议不完整", repair="生成包含 sections 与 renderer 的严格规格")
    charts = [section.get("chart") for section in sections if isinstance(section, dict) and section.get("chart")]
    if not charts:
        _error(issues, "dashboard_without_chart", "交互报告没有可交互图表", repair="连接数据集并生成指标或图表视图；无数据时改用 PPT、Word 或 PDF")


def _evaluate_mermaid(content, issues, checks, metrics):
    source = content.decode("utf-8").strip()
    checks["syntax_header"] = bool(re.match(r"^(?:---[\s\S]*?---\s*)?(flowchart|graph|sequenceDiagram|stateDiagram-v2|erDiagram|timeline|gantt)\b", source, re.IGNORECASE))
    metrics.update({"lines": len(source.splitlines()), "nodes_estimated": len(set(re.findall(r"\b([A-Za-z][A-Za-z0-9_]*)\s*[\[{(]", source)))})
    if not checks["syntax_header"]:
        _error(issues, "mermaid_header", "Mermaid 图缺少有效图类型声明", repair="先选择图类型，再生成对应语法")
    if metrics["nodes_estimated"] > 30:
        _warning(issues, "diagram_complexity", "单图节点过多，难以阅读", repair="拆分总览图和子图")
    if "accTitle" not in source or "accDescr" not in source:
        _warning(issues, "diagram_accessibility", "图缺少可访问标题或描述", repair="补充 accTitle 和 accDescr")


def _evaluate_excalidraw(content, issues, checks, metrics):
    payload = json.loads(content)
    elements = payload.get("elements") if isinstance(payload.get("elements"), list) else []
    live = [item for item in elements if isinstance(item, dict) and not item.get("isDeleted")]
    metrics["element_count"] = len(live)
    checks["schema_valid"] = payload.get("type") == "excalidraw" and bool(live)
    if not checks["schema_valid"]:
        _error(issues, "excalidraw_schema", "白板协议无效或没有元素", repair="根据 WhiteboardSpec 重新渲染")
        return
    ids = {item.get("id") for item in live}
    broken = []
    for item in live:
        for key in ("startBinding", "endBinding"):
            binding = item.get(key)
            if isinstance(binding, dict) and binding.get("elementId") not in ids:
                broken.append(item.get("id"))
    if broken:
        _error(issues, "broken_bindings", "白板存在失效连接", repair="重新绑定连接线端点")
    max_x = max(float(item.get("x", 0)) + abs(float(item.get("width", 0))) for item in live)
    max_y = max(float(item.get("y", 0)) + abs(float(item.get("height", 0))) for item in live)
    metrics.update({"canvas_width": round(max_x), "canvas_height": round(max_y)})
    if max_x > 10000 or max_y > 10000:
        _warning(issues, "canvas_bounds", "白板内容分布过散", repair="自动重新排布并聚焦全部内容")


def _has_content(request):
    return any(section.core_message or section.paragraphs or section.bullets or section.chart for section in request.sections)


def _normalized_text(value: str) -> str:
    return re.sub(r"\s+", "", value).strip("。；;:：").lower()


def _error(issues, code, message, location="", repair="", detail=""):
    issues.append(QualityIssue(code=code, severity="error", message=message + ("：" + detail if detail else ""), location=location, repair_hint=repair))


def _warning(issues, code, message, location="", repair=""):
    issues.append(QualityIssue(code=code, severity="warning", message=message, location=location, repair_hint=repair))
