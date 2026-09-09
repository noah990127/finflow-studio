import asyncio
import io
import json
from pathlib import Path

from docx import Document
from pptx import Presentation
from pypdf import PdfReader

from ..artifacts.planning import prepare_for_format
from ..artifacts.quality import evaluate_artifact
from ..deliverables import (
    create_docx,
    create_excalidraw,
    create_financial_report,
    create_html_slides,
    create_mermaid,
    create_pdf,
    create_pptx,
)
from ..models import DeliverableRequest


def _reference(ref_id: str, name: str, text: str) -> dict:
    return {
        "ref_id": ref_id,
        "resource_id": "eval-" + ref_id,
        "version": 1,
        "source_name": name,
        "text": text,
        "location": {"section": "评估样本"},
        "content_hash": "eval-hash-" + ref_id,
    }


def scenarios() -> dict[str, DeliverableRequest]:
    technology = DeliverableRequest.model_validate({
        "title": "请比较头部科技公司的 2026 战略与当前经营状态，并给出管理层判断",
        "subtitle": "重点关注 AI 基础设施、资本开支、增长质量和主要风险",
        "sections": [{
            "heading": "已核验资料",
            "paragraphs": [
                "公司甲 2026 财年收入同比增长 21%，数据中心业务是主要增量；资本开支同比增长 38%。",
                "公司乙收入同比增长 13%，云业务增速回升，但自由现金流率下降 2.4 个百分点。",
                "共同战略主题是扩大 AI 算力供给、提高开发者粘性并将模型能力嵌入企业软件。",
            ],
            "bullets": ["比较增长质量而非只比较收入增速", "区分已披露事实、计算结果和管理层判断"],
            "refs": [
                _reference("tech-a", "公司甲 2026 财年年报", "收入同比增长 21%，数据中心收入增长 54%，资本开支增长 38%。"),
                _reference("tech-b", "公司乙 2026 财年业绩材料", "收入同比增长 13%，云业务增长 26%，自由现金流率下降 2.4 个百分点。"),
            ],
            "chart": {
                "type": "bar", "title": "收入与资本开支增速对比",
                "categories": ["公司甲收入", "公司甲资本开支", "公司乙收入", "公司乙云业务"],
                "series": [{"name": "同比增速（%）", "values": [21, 38, 13, 26]}],
                "source_ref": "[Ref 1][Ref 2]",
            },
        }],
    })
    operations = DeliverableRequest.model_validate({
        "title": "区域经营质量诊断与回款改善方案",
        "subtitle": "形成可供经营会议审议的正式报告",
        "sections": [{
            "heading": "经营事实",
            "paragraphs": [
                "华南区域本季度收入 1.51 亿元，同比增长 18%，完成预算的 103%。",
                "应收账款周转天数由 42 天上升至 51 天，超过公司 45 天的管理目标。",
                "回款放缓主要集中在五个大型企业客户，相关余额占区域应收的 47%。",
            ],
            "bullets": ["增长与现金转化出现背离", "建议对五个大客户建立周度责任清单"],
            "refs": [_reference("ops", "区域经营台账 2026Q3", "收入、预算、应收余额和客户回款记录已经财务复核。")],
        }, {
            "heading": "行动要求",
            "paragraphs": ["未来六周按客户逐项确认付款节点，销售负责人和财务伙伴共同更新风险状态。"],
            "bullets": ["第 1 周完成余额核对", "第 2 至 4 周解决争议项", "第 6 周复盘现金转化率"],
            "refs": [_reference("policy", "应收账款管理办法", "超过 45 天的应收项目进入重点跟踪清单。")],
        }],
    })
    process = DeliverableRequest.model_validate({
        "title": "请把采购申请到付款的审批过程画清楚，必须包含驳回和补充材料路径",
        "subtitle": "用于新员工理解流程",
        "sections": [{
            "heading": "流程规则",
            "paragraphs": [
                "申请人提交采购用途、预算和供应商材料。部门负责人先审业务必要性，驳回时返回申请人修改。",
                "金额超过十万元时增加财务复核；材料不全时返回申请人补充，复核通过后进入采购执行。",
                "收货验收通过后提交发票，财务核对合同、验收和发票，匹配后安排付款。",
            ],
            "refs": [_reference("procurement", "采购与付款制度", "制度规定了申请、审批、采购、验收、三单匹配和付款步骤。")],
        }],
    })
    dashboard = DeliverableRequest.model_validate({
        "title": "月度销售与毛利交互报告",
        "subtitle": "支持管理层识别增长与利润背离",
        "sections": [{
            "heading": "收入保持增长但毛利率承压",
            "core_message": "8 月收入环比增长 9.2%，毛利率下降 1.6 个百分点",
            "paragraphs": ["收入增长主要来自华东大客户，低毛利项目占比上升导致整体毛利率下降。"],
            "refs": [_reference("sales", "月度销售明细.csv", "包含月份、区域、客户、收入和毛利字段。")],
            "chart": {
                "type": "line", "title": "收入与毛利趋势", "categories": ["6月", "7月", "8月"],
                "series": [
                    {"name": "收入（万元）", "values": [8200, 8600, 9391]},
                    {"name": "毛利（万元）", "values": [2460, 2494, 2573]},
                ],
                "source_ref": "[Ref 1]",
            },
        }],
    })
    return {"technology": technology, "operations": operations, "process": process, "dashboard": dashboard}


async def run(output_directory: Path) -> dict[str, object]:
    output_directory.mkdir(parents=True, exist_ok=True)
    source = scenarios()
    presentation, _ = await prepare_for_format("pptx", source["technology"])
    document, _ = await prepare_for_format("docx", source["operations"])
    diagram, _ = await prepare_for_format("mermaid", source["process"])
    prepared = {
        "pptx": (presentation, create_pptx(presentation), "科技公司战略比较.pptx"),
        "html_slides": (presentation, create_html_slides(presentation), "科技公司战略比较.html"),
        "docx": (document, create_docx(document), "区域经营质量诊断.docx"),
        "pdf": (document, create_pdf(document), "区域经营质量诊断.pdf"),
        "financial_report": (source["dashboard"], create_financial_report(source["dashboard"]), "月度经营交互报告.json"),
        "mermaid": (diagram, create_mermaid(diagram), "采购审批流程.mmd"),
        "excalidraw": (diagram, create_excalidraw(diagram), "采购审批白板.excalidraw"),
    }
    report: dict[str, object] = {}
    for output_format, (request, content, filename) in prepared.items():
        path = output_directory / filename
        path.write_bytes(content)
        quality = evaluate_artifact(output_format, request, content)
        report[output_format] = {
            "file": str(path),
            "quality": quality.model_dump(mode="json"),
            "content": _content_summary(output_format, content),
        }
    (output_directory / "evaluation-report.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return report


def _content_summary(output_format: str, content: bytes) -> dict[str, object]:
    if output_format == "pptx":
        deck = Presentation(io.BytesIO(content))
        titles = []
        for index, slide in enumerate(deck.slides):
            texts = [shape.text.strip() for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()]
            if texts:
                titles.append(texts[1] if index > 0 and len(texts) > 1 and texts[0] != "参考文献" else texts[0])
        return {"pages": len(deck.slides), "page_titles": titles}
    if output_format == "docx":
        document = Document(io.BytesIO(content))
        headings = [p.text for p in document.paragraphs if p.style.name.startswith("Heading")]
        return {"headings": headings, "text": " ".join(p.text for p in document.paragraphs)[:1200]}
    if output_format == "pdf":
        reader = PdfReader(io.BytesIO(content))
        return {"pages": len(reader.pages), "text": " ".join((p.extract_text() or "") for p in reader.pages)[:1200]}
    if output_format in {"financial_report", "excalidraw"}:
        payload = json.loads(content)
        return {"keys": list(payload), "views": len(payload.get("views", [])), "elements": len(payload.get("elements", []))}
    text = content.decode("utf-8")
    if output_format == "html_slides":
        return {"pages": text.count('<section class="slide'), "contains_technical_label": "非 PowerPoint 文件" in text}
    return {"lines": len(text.splitlines()), "preview": text[:1200]}


if __name__ == "__main__":
    asyncio.run(run(Path("/tmp/finflow-result-node-eval")))
